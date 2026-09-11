"""Workbook I/O: reading attachments, filling result templates, checking result contracts."""

from __future__ import annotations

import math
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

BIG_SHEET_CELLS = 300_000


def read_sheets(path: Path, header: int | None = 0) -> dict[str, pd.DataFrame]:
    with pd.ExcelFile(path) as xf:
        return {name: xf.parse(name, header=header) for name in xf.sheet_names}


def sheet_headers(path: Path) -> dict[str, list[Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, list[Any]] = {}
    for ws in wb.worksheets:
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        out[ws.title] = list(row)
    wb.close()
    return out


def _clean(value: Any, decimals: int | None) -> Any:
    if hasattr(value, "item") and not isinstance(value, str | bytes):
        try:
            value = value.item()
        except (ValueError, TypeError):
            pass
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return round(value, decimals) if decimals is not None else value
    return value


def write_result(
    template: Path,
    out: Path,
    sheets: dict[str, tuple[Sequence[Any], Iterable[Sequence[Any]]]],
    *,
    decimals: int | None = 4,
    float_format: str = "0.0000",
) -> dict[str, Any]:
    """Fill a competition result template.

    ``sheets`` maps sheet name -> (header_row, rows). The template's A1 text is always kept;
    header cells after A1 are taken from ``header_row[1:]``. Sheets not mentioned keep their
    template content. Floats are rounded to ``decimals`` and formatted with ``float_format``.
    Very large sheets fall back to openpyxl's write-only mode (no per-cell number format).
    """
    template_wb = openpyxl.load_workbook(template)
    names = list(template_wb.sheetnames)
    unknown = [n for n in sheets if n not in names]
    if unknown:
        raise KeyError(f"sheets {unknown} not in template {template.name} (has {names})")
    a1 = {ws.title: ws.cell(row=1, column=1).value for ws in template_wb.worksheets}
    materialised = {name: (list(header), [list(r) for r in rows]) for name, (header, rows) in sheets.items()}
    cells = sum(max(len(h), 1) * (len(rows) + 1) for h, rows in materialised.values())
    stats: dict[str, Any] = {}
    out.parent.mkdir(parents=True, exist_ok=True)

    if cells > BIG_SHEET_CELLS:
        wb = openpyxl.Workbook(write_only=True)
        for name in names:
            ws = wb.create_sheet(name)
            if name in materialised:
                header, rows = materialised[name]
                full_header = [a1[name]] + list(header[1:])
                ws.append(full_header)
                for row in rows:
                    ws.append([_clean(v, decimals) for v in row])
                stats[name] = {"rows": len(rows), "cols": len(full_header)}
            else:
                for row in template_wb[name].iter_rows(values_only=True):
                    ws.append(list(row))
        wb.save(out)
    else:
        wb = template_wb
        for name, (header, rows) in materialised.items():
            ws = wb[name]
            for row in ws.iter_rows():
                for cell in row:
                    cell.value = None
            full_header = [a1[name]] + list(header[1:])
            for j, value in enumerate(full_header, start=1):
                ws.cell(row=1, column=j, value=_clean(value, None))
            for i, row in enumerate(rows, start=2):
                for j, value in enumerate(row, start=1):
                    cleaned = _clean(value, decimals)
                    cell = ws.cell(row=i, column=j, value=cleaned)
                    if isinstance(cleaned, float) and j > 1:
                        cell.number_format = float_format
            stats[name] = {"rows": len(rows), "cols": len(full_header)}
        wb.save(out)
    return {"path": str(out), "sheets": stats, "write_only": cells > BIG_SHEET_CELLS, "cells": cells}


@dataclass(frozen=True)
class SheetContract:
    name: str
    min_rows: int = 1  # data rows (excluding the header)
    max_rows: int | None = None
    header_len: int | None = None  # exact number of header cells, if known
    numeric_from_col: int | None = None  # 0-based; every data cell from this column must be numeric
    allow_blank: bool = False  # blank numeric cells permitted (e.g. moving-boundary tables)
    decimals: int | None = 4  # numeric cells must carry at most this many decimals
    text_columns: tuple[int, ...] = ()  # 0-based columns that must be non-empty text


@dataclass(frozen=True)
class WorkbookContract:
    file: str
    template: str
    sheets: tuple[SheetContract, ...]


def check_workbook(path: Path, contract: WorkbookContract, template: Path) -> dict[str, Any]:
    errors: list[str] = []
    if not path.exists():
        return {"file": contract.file, "ok": False, "errors": ["file missing"], "sheets": {}}
    try:
        tpl_headers = sheet_headers(template)
    except Exception as exc:  # noqa: BLE001
        tpl_headers = {}
        errors.append(f"template unreadable: {exc!r}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    report: dict[str, Any] = {}
    names = wb.sheetnames
    if tpl_headers and names != list(tpl_headers):
        errors.append(f"sheet names {names} differ from template {list(tpl_headers)}")
    for sc in contract.sheets:
        if sc.name not in names:
            errors.append(f"sheet {sc.name!r} missing")
            continue
        ws = wb[sc.name]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows, ()))
        while header and header[-1] is None:
            header.pop()
        tpl = tpl_headers.get(sc.name)
        if tpl and header and header[0] != tpl[0]:
            errors.append(f"{sc.name}: A1 {header[0]!r} differs from template {tpl[0]!r}")
        if sc.header_len is not None and len(header) != sc.header_len:
            errors.append(f"{sc.name}: header has {len(header)} cells, expected {sc.header_len}")
        n = 0
        bad_numeric = blank = bad_decimals = bad_text = 0
        for row in rows:
            values = list(row)
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            n += 1
            if sc.numeric_from_col is not None:
                for v in values[sc.numeric_from_col : len(header) or None]:
                    if v is None or (isinstance(v, str) and not v.strip()):
                        blank += 1
                    elif not isinstance(v, int | float) or isinstance(v, bool):
                        bad_numeric += 1
                    elif sc.decimals is not None and isinstance(v, float) and abs(round(v, sc.decimals) - v) > 1e-12:
                        bad_decimals += 1
            for j in sc.text_columns:
                v = values[j] if j < len(values) else None
                if v is None or not str(v).strip():
                    bad_text += 1
        if n < sc.min_rows:
            errors.append(f"{sc.name}: {n} data rows < {sc.min_rows}")
        if sc.max_rows is not None and n > sc.max_rows:
            errors.append(f"{sc.name}: {n} data rows > {sc.max_rows}")
        if bad_numeric:
            errors.append(f"{sc.name}: {bad_numeric} non-numeric cells in numeric region")
        if blank and not sc.allow_blank:
            errors.append(f"{sc.name}: {blank} blank cells in numeric region")
        if bad_decimals:
            errors.append(f"{sc.name}: {bad_decimals} cells with more than {sc.decimals} decimals")
        if bad_text:
            errors.append(f"{sc.name}: {bad_text} empty cells in required text columns")
        report[sc.name] = {"rows": n, "header": [str(h) for h in header[:6]] + (["…"] if len(header) > 6 else []),
                           "blank": blank}
    wb.close()
    return {"file": contract.file, "ok": not errors, "errors": errors, "sheets": report}
